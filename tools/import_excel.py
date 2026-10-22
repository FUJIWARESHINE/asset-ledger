#!/usr/bin/env python3
"""把「虚拟资产.xlsx」转换成 asset-ledger 的导入 JSON。

一个域名在这份表里有两个维度，脚本两者都保留：

  · 申请商  provider / provider_account  —— 在 dpdns / dnshe 上申请，管到期与配额
  · 托管方  platform / account            —— NS 指到哪个 Cloudflare 账号，管 DNS 与路由

托管方是资产的归属主体（落成 account_id），所以域名挂在 Cloudflare 账号下；
申请商落成资产上的两个独立字段，卡片与磁贴上用徽章显示。
没托管的域名（cc.cd / us.ci / cn.mt 这类只申请没上线）挂到申请商账号下。

用法（需要 openpyxl：`.venv/Scripts/pip install openpyxl`）：
    python tools/import_excel.py                     # 只打印报告
    python tools/import_excel.py --write             # 额外写出导入 JSON
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = Path("虚拟资产.xlsx")
DEFAULT_OUT = BASE_DIR / "data" / "import-virtual-assets.json"

CF = "Cloudflare"
# 后缀 → 申请商。dpdns 提供 .dpdns.org 与 .qzz.io；dnshe 提供下面这批；
# eu.org 是独立的老牌免费域名服务，不属于这两家。
SUFFIX_PROVIDER = [
    (".dpdns.org", "dpdns"),
    (".qzz.io", "dpdns"),
    (".de5.net", "dnshe"),
    (".cc.cd", "dnshe"),
    (".us.ci", "dnshe"),
    (".cn.mt", "dnshe"),
    (".bot.cd", "dnshe"),
    (".bbroot.com", "dnshe"),
    (".eu.org", "eu.org"),
]
SERVICE_KIND = {"pages": "pages", "workers": "worker", "worker": "worker"}


def clean(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\u3000", " ")).strip()


def merge_map(ws) -> dict[tuple[int, int], tuple[int, int]]:
    """(row, col) -> 合并区域左上角 (row, col)，用来做「语义分组合并」的展开。"""
    out: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = (rng.min_row, rng.min_col)
    return out


def at(ws, mm, row: int, col: int) -> str:
    r0, c0 = mm.get((row, col), (row, col))
    return clean(ws.cell(row=r0, column=c0).value)


def is_block_head(mm, row: int, col: int) -> bool:
    """该单元格是否是一个纵向合并块的首行（服务只在首行产出一条）。"""
    tl = mm.get((row, col))
    return tl is None or tl[0] == row


def provider_of(domain: str) -> str:
    d = domain.lower()
    for suffix, prov in SUFFIX_PROVIDER:
        if d.endswith(suffix):
            return prov
    return ""


def split_service(raw: str) -> tuple[str, str]:
    """'Pages：CFNews' -> ('pages', 'CFNews')"""
    parts = re.split(r"[：:]", raw, maxsplit=1)
    if len(parts) == 2:
        kind, name = parts[0].strip(), parts[1].strip()
        return SERVICE_KIND.get(kind.lower(), "other"), name
    return "other", raw.strip()


class Collector:
    def __init__(self) -> None:
        self.accounts: dict[tuple[str, str], dict] = {}
        self.domains: dict[str, dict] = {}
        self.services: list[dict] = []

    def add_account(self, platform: str, name: str, **kw) -> None:
        if not name:
            return
        key = (platform.lower(), name.lower())
        if key in self.accounts:
            old = self.accounts[key]
            for k, v in kw.items():
                if v and not old.get(k):
                    old[k] = v
            return
        self.accounts[key] = {"platform": platform, "name": name,
                              "label": "", "url": "", "status": "active",
                              "credential_hint": "", "notes": "", **kw}

    def domain(self, name: str) -> dict:
        if name not in self.domains:
            self.domains[name] = {"name": name, "host_account": "", "host_platform": "",
                                  "provider": provider_of(name), "provider_account": "",
                                  "forever": False, "app": "", "note": ""}
        return self.domains[name]


def read_dnshe(ws, mm, col: Collector) -> None:
    """Sheet2：申请商侧的事实源——谁申请的、到期、配额、托管指向。"""
    for r in range(2, ws.max_row + 1):
        acc = at(ws, mm, r, 2)
        dom = at(ws, mm, r, 3)
        if not dom:
            continue
        col.add_account("Dnshe", acc, notes=f"免费域名配额 {at(ws, mm, r, 7)}" if at(ws, mm, r, 7) else "")
        d = col.domain(dom)
        d["provider"] = "dnshe"
        d["provider_account"] = d["provider_account"] or acc
        if at(ws, mm, r, 6).lower().startswith("forever"):
            d["forever"] = True
        host = clean(ws.cell(row=r, column=4).value)
        if host and not d["host_account"]:
            # 'cloudflare：1478289779' / 'cloudflare：科信研测'
            side, _, who = host.replace("：", ":").partition(":")
            d["host_platform"] = side.strip() or CF
            d["host_hint"] = who.strip()
        app = clean(ws.cell(row=r, column=5).value)
        if app:
            d["app"] = app


def read_cloudflare(ws, mm, col: Collector) -> None:
    """Sheet1：托管侧的事实源——域名挂在哪个 CF 账号、上面跑了什么服务。"""
    cf_accounts: list[str] = []
    last_platform = ""
    for r in range(2, ws.max_row + 1):
        platform = at(ws, mm, r, 1) or last_platform
        if not platform:
            continue
        last_platform = platform
        account = at(ws, mm, r, 2)

        if platform.lower() == "cloudflare":
            col.add_account(CF, account, url="https://dash.cloudflare.com")
            if account and account not in cf_accounts:
                cf_accounts.append(account)
            dom_raw = at(ws, mm, r, 3)
            if dom_raw and dom_raw != "/":
                d = col.domain(dom_raw)
                if not d["host_account"]:
                    d["host_account"] = account
                    d["host_platform"] = CF
            # 服务：D 列纵向合并 = 一个服务覆盖一个域名块，只在块首行产出
            srv_raw = at(ws, mm, r, 4)
            if srv_raw and is_block_head(mm, r, 4):
                kind, sname = split_service(srv_raw)
                col.services.append({
                    "name": sname,
                    "category": kind,
                    "platform": CF,
                    "account": account,
                    "domain": at(ws, mm, r, 3),
                    "url": at(ws, mm, r, 5),
                    "deploy": at(ws, mm, r, 6),
                    "note": at(ws, mm, r, 7),
                })
        elif platform.lower() == "github":
            col.add_account("GitHub", account, url="https://github.com")
        elif platform.lower() == "vercel":
            col.add_account("Vercel", account, url="https://vercel.com")
            srv_raw = at(ws, mm, r, 4)
            if srv_raw:
                _, sname = split_service(srv_raw)
                col.services.append({
                    "name": sname, "category": "vercel",
                    "platform": "Vercel", "account": account,
                    "domain": "", "url": at(ws, mm, r, 5),
                    "deploy": at(ws, mm, r, 6), "note": at(ws, mm, r, 7),
                })
        elif platform.lower() == "dpdns":
            col.add_account("Dpdns", account, notes="免费域名商")
    # dpdns 后缀的域名，申请账号与托管账号是同一批邮箱（表里只列了部分账号，这里补齐）
    for d in col.domains.values():
        if d["provider"] == "dpdns":
            d["provider_account"] = d["provider_account"] or d["host_account"]
            if d["provider_account"]:
                col.add_account("Dpdns", d["provider_account"], notes="免费域名商")
        elif d["provider"] and not d["provider_account"]:
            d["provider_account"] = d["host_account"]


def resolve_hosts(col: Collector) -> None:
    """Sheet2 的托管列用的是邮箱前缀或账号别名，这里对齐到 Sheet1 的真实账号名。"""
    cf_names = [a["name"] for a in col.accounts.values() if a["platform"] == CF]
    for d in col.domains.values():
        hint = d.pop("host_hint", "")
        if not hint or d["host_account"]:
            continue
        hit = ""
        for name in cf_names:
            local = name.split("@")[0]
            if hint == name or hint == local or hint in name:
                hit = name
                break
        if hit:
            d["host_account"] = hit
            d["host_platform"] = CF
        else:
            # 表里出现但 Sheet1 没有的 CF 账号（如「科信研测」），补建
            col.add_account(CF, hint, url="https://dash.cloudflare.com")
            d["host_account"] = hint
            d["host_platform"] = CF


def build(col: Collector) -> tuple[list[dict], list[dict]]:
    assets: list[dict] = []

    for d in col.domains.values():
        hosted = bool(d["host_account"])
        if hosted:
            platform = d["host_platform"] or CF
        elif d["provider"] == "dnshe":
            platform = "Dnshe"
        elif d["provider"] == "dpdns":
            platform = "Dpdns"
        else:
            platform = "未归类"
        account = d["host_account"] or d["provider_account"]
        mem = ["免费域名"]
        if d["forever"]:
            mem.append("永久有效")
        if d["app"]:
            mem.append(f"应用：{d['app']}")
        if d["note"]:
            mem.append(d["note"])
        tags = [] if hosted else ["未托管"]
        assets.append({
            "name": d["name"], "category": "domain",
            "platform": platform, "account": account,
            "provider": d["provider"], "provider_account": d["provider_account"],
            "url": "", "environment": "prod", "status": "active",
            "expires_at": "", "auto_renew": False,
            "cost": 0, "currency": "CNY", "cycle": "none",
            "credential_hint": "", "tags": ",".join(tags),
            "notes": " · ".join(mem),
        })

    for s in col.services:
        mem = []
        if s["deploy"]:
            mem.append("自动部署" if s["deploy"].lower() == "auto" else f"部署 {s['deploy']}")
        if s["note"]:
            mem.append(s["note"])
        tags = ["未使用"] if "未使用" in s["note"] else []
        assets.append({
            "name": s["name"], "category": s["category"],
            "platform": s["platform"], "account": s["account"],
            "provider": "", "provider_account": "",
            "url": s["url"], "environment": "prod", "status": "active",
            "expires_at": "", "auto_renew": False,
            "cost": 0, "currency": "CNY", "cycle": "none",
            "credential_hint": "", "tags": ",".join(tags),
            "notes": " · ".join(mem),
        })

    accounts = list(col.accounts.values())
    accounts.sort(key=lambda a: (a["platform"], a["name"]))
    assets.sort(key=lambda a: (a["platform"], a["account"], a["category"], a["name"]))
    return accounts, assets


def report(accounts: list[dict], assets: list[dict]) -> None:
    print(f"账号 {len(accounts)} 个 / 资产 {len(assets)} 个\n")
    by_plat: dict[str, list[dict]] = {}
    for a in accounts:
        by_plat.setdefault(a["platform"], []).append(a)
    for p, accs in by_plat.items():
        n = sum(1 for x in assets if x["platform"] == p)
        print(f"[{p}] {len(accs)} 账号 / {n} 资产")
        for a in accs:
            mine = [x for x in assets if x["platform"] == p and x["account"] == a["name"]]
            if not mine:
                print(f"    {a['name']:<28} —")
                continue
            print(f"    {a['name']:<28} {len(mine)} 项")
            for x in mine:
                if x["category"] == "domain":
                    tag = f"申请 {x['provider'] or '?':<7}"
                    print(f"        {x['name']:<24} {tag} {'托管 ✓' if x['platform'] == CF else '未托管'}")
                else:
                    print(f"        {x['name']:<24} [{x['category']}] {x['url']}")
        print()

    orphans = [x for x in assets if not x["platform"]]
    if orphans:
        print(f"⚠ 未归属资产 {len(orphans)} 个")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    ap.add_argument("--write", action="store_true")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    col = Collector()
    ws2 = wb["Dnshe"]
    read_dnshe(ws2, merge_map(ws2), col)
    ws1 = wb["Cloudflare"]
    read_cloudflare(ws1, merge_map(ws1), col)
    resolve_hosts(col)

    accounts, assets = build(col)
    report(accounts, assets)

    if args.write:
        payload = {"mode": "replace", "accounts": accounts, "assets": assets}
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"已写出 {out}  ({out.stat().st_size} 字节)")


if __name__ == "__main__":
    main()
